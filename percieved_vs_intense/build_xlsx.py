"""
Stevens's power law: perceived sensation against physical intensity.

Every curve is the same one-line model,  psi = I^a , drawn for six different
values of the exponent a. The exponent is the whole story:

    a > 1   the sensation grows faster than the stimulus. Double the current
            and the shock feels far more than twice as bad.
    a = 1   veridical. What you encode is what the reader perceives.
    a < 1   the sensation is compressed. Double the area and it looks like
            rather less than twice as much.

Which is why this chart belongs next to the effectiveness ranking. Length sits
at a = 1 and is the only channel here that a reader reads back faithfully; area
at 0.7 is systematically under-read, so a bubble twice the area of another does
not look twice as big. That is not a preference — it is a measured bias.

The workbook is built so a reader can change an exponent in row 7 and watch its
curve move: the table is live formulas over that row, and the chart reads the
table.

Run:  python build_xlsx.py
"""

import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.chart import Reference, ScatterChart, Series
from openpyxl.chart.marker import Marker
from openpyxl.drawing.line import LineProperties
from openpyxl.chart.shapes import GraphicalProperties

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE / "style"))
import xl_style as xs                                    # noqa: E402

OUT = HERE / "perceived_vs_intensity.xlsx"

# --- the data -----------------------------------------------------------
# name, exponent, colour, and where to put the series' own label.
#
# Five of the six exponents are Stevens's own published values (Stevens 1957,
# Table 1). DEPTH IS NOT: 0.67 is Stevens's exponent for LOUDNESS, and his table
# carries no depth entry at all. It is kept here because it is the figure this
# reproduces and because depth's rank is well established in the visualization
# literature — but the sheet says so rather than passing it off as Stevens.
CHANNELS = [
    # name                 a      colour    label idx   pos
    ("Electric Shock",    3.5,  "1B9AD6",        30,   "l"),
    ("Saturation",        1.7,  "2E9E4F",        47,   "l"),
    ("Length",            1.0,  "7B3F9E",       100,   "l"),
    ("Area",              0.7,  "C0392B",        90,   "t"),
    ("Depth",            0.67,  "0D7680",        78,   "b"),
    ("Brightness",        0.5,  "8C8C8C",        95,   "b"),
]

X_MAX, X_STEP = 5.0, 0.05
N = int(round(X_MAX / X_STEP)) + 1                       # 101 points

# --- sheet geometry -----------------------------------------------------
LABEL_COL = 2                     # B — the label gutter
TBL_COL = 3                       # C — intensity
EXP_ROW = 7                       # the exponents the formulas read
HDR_ROW = 8
TOP = 9                           # first data row
BOT = TOP + N - 1

Y_MAX, Y_UNIT = 5.4, 1.0          # 5.4 leaves the "5" tick with room above it
PANEL_W, PANEL_H = 19.0, 12.5



def vertical_axis_title(axis, text, size_pt=11, color=xs.TEXT_2):
    """
    A value-axis title turned on its side, reading bottom-to-top.

    xl_style.axis_title deliberately forces rot=0 — the house guide forbids
    rotated axis titles, and that rule stands for the rest of the workbooks, so
    this is a local override rather than an edit to the shared module.

    rot is in 60000ths of a degree: -5400000 is -90.
    """
    cp = xs._cp(size_pt, color)
    para = xs.Paragraph(pPr=xs.ParagraphProperties(algn="ctr", defRPr=cp),
                        r=[xs.RegularTextRun(rPr=cp, t=text)])
    axis.title = xs.Title(
        tx=xs.Text(rich=xs.RichText(
            bodyPr=xs.RichTextProperties(rot=-5400000, vert="horz"),
            p=[para])),
        overlay=False)
    return axis


def build():
    wb = Workbook()
    ws = wb.active
    ws.title = "Stevens_Law"

    xs.write_header(
        ws, 2,
        "Only length is perceived in proportion to what it encodes",
        "Perceived sensation against physical intensity, psi = I^a, for six channels")

    xs.note(ws, 5,
            "Change an exponent in row 7 and its curve moves — the table is live "
            "formulas over that row, and the chart reads the table.")

    # --- exponent row, which the whole table hangs off ------------------
    # NB: styles.Font here, not xs._cp — _cp builds a drawing
    # CharacterProperties for chart rich text. Assigned to cell.font it
    # writes XML Excel refuses to open, with no complaint from Python.
    EXP_FONT = Font(name=xs.FONT, size=10, bold=True, color=xs.TEXT_2)
    c = ws.cell(row=EXP_ROW, column=LABEL_COL, value="Exponent (a)")
    c.font = EXP_FONT
    for j, (_, a, _, _, _) in enumerate(CHANNELS):
        c = ws.cell(row=EXP_ROW, column=TBL_COL + 1 + j, value=a)
        c.font = EXP_FONT
        c.number_format = "0.##"

    # --- table header ---------------------------------------------------
    # The display name carries the exponent too, because the chart's direct
    # labels are drawn from these cells and the reader needs the number there.
    headers = ["Intensity (I)"] + [f"{n} ({a:g})" for n, a, _, _, _ in CHANNELS]
    xs.write_table_header(ws, HDR_ROW, headers, col=TBL_COL)

    # --- the table ------------------------------------------------------
    for i in range(N):
        r = TOP + i
        x = round(i * X_STEP, 4)
        cx = ws.cell(row=r, column=TBL_COL, value=x)
        cx.number_format = "0.00"
        for j in range(len(CHANNELS)):
            col = TBL_COL + 1 + j
            letter = ws.cell(row=EXP_ROW, column=col).column_letter
            cell = ws.cell(row=r, column=col,
                           value=f"=${ws.cell(row=r, column=TBL_COL).column_letter}"
                                 f"{r}^{letter}${EXP_ROW}")
            cell.number_format = "0.00"

    src_row = BOT + 2
    xs.write_source(
        ws, src_row,
        "Exponents: Stevens, On the psychophysical law, Psychological Review 64(3), "
        "1957 — electric shock 3.5, redness/saturation 1.7, visual length 1.0, "
        "visual area 0.7, brightness 0.5 (point source). Depth 0.67 is NOT in "
        "Stevens's table (0.67 is his loudness exponent); it is carried from the "
        "visualization literature, where depth ranks just below area.")

    xs.widths(ws, [2, 16] + [15] * (len(CHANNELS) + 1), start=1)

    # --- the chart ------------------------------------------------------
    chart = ScatterChart()
    chart.scatterStyle = "lineMarker"
    chart.title = xs.chart_message(
        "Only length is read back faithfully — everything else is biased",
        size_pt=13)

    xref = Reference(ws, min_col=TBL_COL, min_row=TOP, max_row=BOT)
    for j, (name, a, colour, lbl_idx, pos) in enumerate(CHANNELS):
        col = TBL_COL + 1 + j
        s = Series(Reference(ws, min_col=col, min_row=HDR_ROW, max_row=BOT),
                   xref, title_from_data=True)
        s.marker = Marker(symbol="none")
        s.smooth = True
        s.graphicalProperties = GraphicalProperties(
            ln=LineProperties(solidFill=colour, w=xs.W_HIGHLIGHT))
        xs.label_point(s, lbl_idx, colour, size=10, position=pos)
        chart.series.append(s)

    # No gridlines: the chart is read for the SHAPE of each curve, not for
    # values off it, and the reference figure has none either.
    xs.declutter(chart, gridlines=False)
    xs.no_legend(chart)                      # every curve is labelled on itself

    chart.x_axis.scaling.min, chart.x_axis.scaling.max = 0, X_MAX
    chart.x_axis.majorUnit = 1
    chart.x_axis.numFmt = "0"
    chart.y_axis.scaling.min, chart.y_axis.scaling.max = 0, Y_MAX
    chart.y_axis.majorUnit = Y_UNIT
    chart.y_axis.numFmt = "0"
    # Both axes are drawn on the reference; declutter only restores the x one.
    chart.y_axis.spPr = GraphicalProperties(
        ln=LineProperties(solidFill=xs.TEXT_2, w=xs.W_THIN))
    xs.axis_title(chart.x_axis, "Physical intensity", size_pt=11)
    vertical_axis_title(chart.y_axis, "Perceived sensation")

    chart.width, chart.height = PANEL_W, PANEL_H
    ws.add_chart(chart, "K7")

    ws.sheet_view.showGridLines = False
    wb.save(OUT)
    print(f"wrote {OUT.name}  ({N} points, {len(CHANNELS)} channels)")


if __name__ == "__main__":
    build()
