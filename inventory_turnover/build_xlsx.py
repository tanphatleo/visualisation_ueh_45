"""
Build inventory_turnover.xlsx — two renderings of one dataset.

    sheet "replica"     a pixel-faithful rebuild of image.png (palette and geometry
                        sampled straight out of the PNG)
    sheet "standup"     the same data under style_chart.md v1.0, "Warehouse Stand-Up":
                        cream paper, tangerine hero line, punchline title, an event
                        line at the year the story turns

Both charts read the table on their own sheet, and the tables come from
data/inventory_turnover.csv. Change a cell, the chart moves.

    C:\\Python313\\python.exe build_xlsx.py
"""

from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import LineChart, Reference, ScatterChart, Series
from openpyxl.chart.data_source import StrRef
from openpyxl.chart.label import DataLabel, DataLabelList
from openpyxl.chart.layout import Layout, ManualLayout
from openpyxl.chart.marker import Marker
from openpyxl.chart.series import SeriesLabel
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.chart.text import RichText, Text
from openpyxl.chart.title import Title
from openpyxl.drawing.line import LineProperties
from openpyxl.drawing.text import (
    CharacterProperties,
    Font as DFont,
    Paragraph,
    ParagraphProperties,
    RegularTextRun,
    RichTextProperties,
)
from openpyxl.styles import Alignment, Font, PatternFill

HERE = Path(__file__).parent
CSV = HERE / "data" / "inventory_turnover.csv"
XLSX = HERE / "inventory_turnover.xlsx"

# ---------------------------------------------------------------------------
# palette 1 — sampled out of image.png, for the replica sheet only
# ---------------------------------------------------------------------------
R_BLUE = "357DBF"        # OUR COMPANY line + its end label
R_DARK = "404040"        # INDUSTRY BENCHMARK dashed line, titles
R_TICK = "595959"
R_AXIS = "B5B5B5"
R_FONT = "Segoe UI"

# ---------------------------------------------------------------------------
# palette 2 — style_chart.md v1.0 "Warehouse Stand-Up" §2.1 / §9
# ---------------------------------------------------------------------------
SURFACE, WASH, GRID = "FFF9F0", "FBEFDD", "E7D8C3"
CONTEXT, TEXT_2, INK = "6E6259", "5C5149", "2F2A26"
ACCENT, ACCENT_TEXT, BAD = "EF5B23", "C2400F", "B8352C"
FONT_HEAD, FONT_BODY = "Verdana", "Segoe UI"
TURNS_FMT = '0.0"×"'

PT = 12700
W_HERO, W_CONTEXT, W_AXIS, W_EVENT = int(3 * PT), int(1.5 * PT), int(1 * PT), int(1 * PT)
CHART_W, CHART_H = 14.6, 10.9            # cm, 4:3 — §4

PUNCHLINE = "Four years of losing to the Joneses — then someone found the warehouse keys."
SUBTITLE = ("Inventory turns per year, fiscal years ending 12/31. Higher is better; "
            "the dashed line is the industry benchmark.")
SOURCE = ("Source: data/inventory_turnover.csv  |  Turns = COGS ÷ average inventory. "
          "No inventory was harmed in the making of this chart.")
EVENT = (2024, "New WMS goes live 🎉")   # the chart's one emoji and one wisecrack


# ---------------------------------------------------------------------------
# text helpers — no raw hex or magic size at a call site
# ---------------------------------------------------------------------------
def _cp(size_pt, color, bold=False, font=FONT_BODY):
    return CharacterProperties(sz=int(size_pt * 100), b=bold, solidFill=color,
                               latin=DFont(typeface=font))


def _rich(size_pt, color, bold=False, rot=None, font=FONT_BODY):
    cp = _cp(size_pt, color, bold, font)
    body = RichTextProperties(rot=rot, vert="horz") if rot else RichTextProperties()
    return RichText(bodyPr=body,
                    p=[Paragraph(pPr=ParagraphProperties(defRPr=cp), endParaRPr=cp)])


def _title(text, size_pt, color, bold=False, rot=None, left=None, top=None, font=FONT_BODY):
    """A chart/axis title. `left`/`top` pin it as a fraction of the chart area."""
    cp = _cp(size_pt, color, bold, font)
    body = RichTextProperties(rot=rot, vert="horz") if rot else RichTextProperties()
    # algn="l": without it a title that wraps comes out centred, and the second
    # line sits indented under the first — §1, top-left carries the message.
    para = Paragraph(pPr=ParagraphProperties(defRPr=cp, algn="l"),
                     r=[RegularTextRun(rPr=cp, t=text)])
    t = Title(tx=Text(rich=RichText(bodyPr=body, p=[para])), overlay=False)
    if left is not None:
        t.layout = Layout(manualLayout=ManualLayout(
            xMode="edge", yMode="edge", x=left, y=top))
    return t


def _end_label(series, idx, color, bold=True, size_pt=9):
    """
    Direct labelling: the series NAME printed at its last point.

    This is what replaces the legend — the words sit next to the line they
    describe, so the reader never matches a colour to a key.
    """
    off = dict(showVal=False, showCatName=False, showLegendKey=False,
               showPercent=False, showBubbleSize=False)
    lbl = DataLabel(idx=idx, showSerName=True, dLblPos="r", **off)
    lbl.txPr = _rich(size_pt, color, bold=bold)
    series.dLbls = DataLabelList(dLbl=[lbl], showSerName=False, **off)
    return series


def _axis(ax, line_color, tick_color, tick_pt=9, num_fmt="0"):
    ax.delete = False                    # never leave this unset: Excel then draws no labels
    ax.majorTickMark, ax.minorTickMark = "out", "none"
    ax.numFmt = ax.number_format = num_fmt
    ax.txPr = _rich(tick_pt, tick_color)
    ax.spPr = GraphicalProperties(ln=LineProperties(solidFill=line_color, w=W_AXIS))
    return ax


# ---------------------------------------------------------------------------
# sheet 1 — the replica
# ---------------------------------------------------------------------------
def build_replica(ws, df):
    ws.sheet_view.showGridLines = False
    ws.cell(1, 2, "Replica of image.png — palette and geometry sampled from the PNG").font = (
        Font(name=R_FONT, size=10, italic=True, color=R_AXIS))

    head = 3
    for j, h in enumerate(["FISCAL YEAR ENDING 12/31", "OUR COMPANY", "INDUSTRY BENCHMARK"]):
        c = ws.cell(head, 2 + j, h)
        c.font = Font(name=R_FONT, size=10, bold=True, color=R_DARK)
        c.alignment = Alignment(horizontal="center", wrap_text=True)
    ws.row_dimensions[head].height = 30

    for i, row in enumerate(df.itertuples(index=False), start=head + 1):
        ws.cell(i, 2, int(row.fiscal_year)).font = Font(name=R_FONT, size=11, color=R_TICK)
        for j, v in enumerate((row.our_company, row.industry_benchmark)):
            c = ws.cell(i, 3 + j, float(v))
            c.font = Font(name=R_FONT, size=11, color=R_TICK)
            c.number_format = "0.0"
            c.alignment = Alignment(horizontal="center")

    first, last = head + 1, head + len(df)
    ws.cell(last + 2, 2, "Source: data/inventory_turnover.csv — turns = COGS ÷ average "
                         "inventory, fiscal years ending 12/31.").font = Font(
        name=R_FONT, size=9, color=R_AXIS)
    for col, w in zip("BCD", (26, 14, 20)):
        ws.column_dimensions[col].width = w

    chart = LineChart()
    chart.title = _title("Annual inventory turnover", 14, R_DARK, left=0.015, top=0.02)
    chart.add_data(Reference(ws, min_col=3, max_col=4, min_row=head, max_row=last),
                   titles_from_data=True)
    chart.set_categories(Reference(ws, min_col=2, min_row=first, max_row=last))

    chart.width, chart.height = CHART_W, CHART_H     # image.png is 402x301 px ≈ 4:3
    chart.roundedCorners = False
    chart.style = None
    chart.legend = None
    chart.graphical_properties = GraphicalProperties(
        solidFill="FFFFFF", ln=LineProperties(noFill=True))
    # 26% of the width kept free on the right: Excel silently clips a data label
    # that runs past the chart edge.
    chart.layout = Layout(manualLayout=ManualLayout(
        xMode="edge", yMode="edge", x=0.175, y=0.145, w=0.565, h=0.665))

    ours, bench = chart.series
    ours.graphicalProperties = GraphicalProperties(
        ln=LineProperties(solidFill=R_BLUE, w=int(2.25 * PT)))
    ours.marker, ours.smooth = Marker(symbol="none"), False
    _end_label(ours, len(df) - 1, R_BLUE, bold=True)

    bench.graphicalProperties = GraphicalProperties(
        ln=LineProperties(solidFill=R_DARK, w=W_CONTEXT, prstDash="dash"))
    bench.marker, bench.smooth = Marker(symbol="none"), False
    _end_label(bench, len(df) - 1, R_DARK, bold=False)

    y = _axis(chart.y_axis, R_AXIS, R_TICK)
    y.scaling.min, y.scaling.max, y.majorUnit = 0, 10, 2
    y.majorGridlines = None
    y.title = _title("# OF INVENTORY TURNS", 8, R_DARK, rot=-5400000)
    # "On tick marks", not "between": 2020 sits ON the value axis, as in image.png.
    # crossBetween lives on the VALUE axis, not the category one.
    y.crossBetween = "midCat"

    x = _axis(chart.x_axis, R_AXIS, R_TICK)
    x.majorGridlines = None
    x.title = _title("FISCAL YEAR ENDING 12/31", 8, R_DARK, left=0.175, top=0.875)
    x.lblOffset = 100

    ws.add_chart(chart, f"F{head}")


# ---------------------------------------------------------------------------
# sheet 2 — style_chart.md v1.0
# ---------------------------------------------------------------------------
def event_line(ws, row, x_index, caption, n_categories, top=0.95):
    """
    A vertical reference line built INTO the chart — style_chart.md §7.

    Two points at the same x, y = 0 and y = top, drawn as a scatter line on a
    hidden 0-1 secondary axis pair. The caption is the series NAME, shown as the
    top point's data label, so the words travel with the line.
    """
    c = ws.cell(row, 11, caption)
    c.font = Font(name=FONT_BODY, size=9, italic=True, color=BAD)
    for k, y in enumerate((0, top)):
        ws.cell(row + k, 12, x_index).font = Font(name=FONT_BODY, size=9, color=TEXT_2)
        ws.cell(row + k, 13, y).font = Font(name=FONT_BODY, size=9, color=TEXT_2)

    s = Series(Reference(ws, min_col=13, min_row=row, max_row=row + 1),
               Reference(ws, min_col=12, min_row=row, max_row=row + 1))
    # Series(title=...) str()s whatever it is given, so the caption has to be
    # attached afterwards or it serialises as a repr.
    s.tx = SeriesLabel(strRef=StrRef(f"'{ws.title}'!$K${row}"))
    s.graphicalProperties = GraphicalProperties(
        ln=LineProperties(solidFill=BAD, w=W_EVENT, prstDash="dash"))
    s.marker = Marker(symbol="none")

    off = dict(showVal=False, showCatName=False, showLegendKey=False,
               showPercent=False, showBubbleSize=False)
    top_label = DataLabel(idx=1, showSerName=True, dLblPos="t", **off)
    top_label.txPr = _rich(9, BAD, bold=False)
    s.dLbls = DataLabelList(dLbl=[DataLabel(idx=0, showSerName=False, **off), top_label],
                            showSerName=False, **off)

    sc = ScatterChart()
    sc.series.append(s)
    # The secondary pair must sit OPPOSITE the primary pair and cross at max, or
    # Excel rescales the PRIMARY axis to the 0-1 helper values and flattens the data.
    sc.x_axis.axId, sc.y_axis.axId = 300, 200
    sc.x_axis.crossAx, sc.y_axis.crossAx = 200, 300
    sc.x_axis.axPos, sc.y_axis.axPos = "t", "r"
    sc.x_axis.crosses = sc.y_axis.crosses = "max"
    # 1 .. n, not 0,5 .. n+0,5: the value axis uses crossBetween="midCat", so
    # category i sits ON tick i. Get this wrong and the line lands between years.
    sc.x_axis.scaling.min, sc.x_axis.scaling.max = 1, n_categories
    sc.y_axis.scaling.min, sc.y_axis.scaling.max = 0, 1
    sc.x_axis.majorGridlines = sc.y_axis.majorGridlines = None
    for ax in (sc.x_axis, sc.y_axis):
        # Hidden, but NOT deleted: deleting it makes Excel reassign the series to
        # the primary axis, which then rescales to 0-1.
        ax.delete = False
        ax.majorTickMark = ax.minorTickMark = "none"
        # tickLblPos="none" is read by openpyxl as "unset" and dropped; ";;;" is
        # Excel's own show-nothing number format and does the job.
        ax.numFmt = ";;;"
        ax.spPr = GraphicalProperties(ln=LineProperties(noFill=True))
    return sc


def build_standup(ws, df):
    ws.sheet_view.showGridLines = False
    n = len(df)

    # The punchline rides on the chart object so the chart is deck-ready; the
    # cells carry the straight-man subtitle and the source. Never both. (§4)
    c = ws.cell(2, 2, SUBTITLE)
    c.font = Font(name=FONT_BODY, size=11, color=TEXT_2)

    # the table, out of the chart's way
    head = 2
    for j, h in enumerate(["FISCAL YEAR ENDING 12/31", "US", "THE JONESES"]):
        cell = ws.cell(head, 11 + j, h)
        cell.font = Font(name=FONT_BODY, size=10, bold=True, color=INK)
        cell.fill = PatternFill("solid", fgColor=WASH)
        cell.alignment = Alignment(horizontal="center", wrap_text=True)
    ws.row_dimensions[head].height = 30

    for i, row in enumerate(df.itertuples(index=False), start=head + 1):
        ws.cell(i, 11, int(row.fiscal_year)).font = Font(name=FONT_BODY, size=11, color=TEXT_2)
        for j, v in enumerate((row.our_company, row.industry_benchmark)):
            cell = ws.cell(i, 12 + j, float(v))
            cell.font = Font(name=FONT_BODY, size=11, color=INK)
            cell.number_format = TURNS_FMT
            cell.alignment = Alignment(horizontal="center")

    first, last = head + 1, head + n
    for col, w in zip("KLM", (26, 10, 14)):
        ws.column_dimensions[col].width = w

    ev_row = last + 3
    note = ws.cell(ev_row - 1, 11,
                   "Event line helper — two points at the same x, y=0 and y=0,95, drawn as a "
                   "scatter line on a hidden 0–1 axis. The caption is the series NAME. "
                   "The hidden axis pair encodes nothing; do not 'fix' it.")
    note.font = Font(name=FONT_BODY, size=9, italic=True, color=BAD)

    chart = LineChart()
    chart.title = _title(PUNCHLINE, 13, INK, bold=True, left=0.012, top=0.02, font=FONT_HEAD)
    chart.add_data(Reference(ws, min_col=12, max_col=13, min_row=head, max_row=last),
                   titles_from_data=True)
    chart.set_categories(Reference(ws, min_col=11, min_row=first, max_row=last))

    chart.width, chart.height = CHART_W, CHART_H
    chart.roundedCorners = False
    chart.style = None
    chart.legend = None                       # §11: direct labels, never a legend
    chart.graphical_properties = GraphicalProperties(
        solidFill=SURFACE, ln=LineProperties(noFill=True))
    chart.layout = Layout(manualLayout=ManualLayout(
        xMode="edge", yMode="edge", x=0.175, y=0.20, w=0.555, h=0.60))

    ours, bench = chart.series
    ours.graphicalProperties = GraphicalProperties(
        ln=LineProperties(solidFill=ACCENT, w=W_HERO, cap="rnd"))
    # hollow round markers: the reader can count the periods
    ours.marker = Marker(symbol="circle", size=7, spPr=GraphicalProperties(
        solidFill=SURFACE, ln=LineProperties(solidFill=ACCENT, w=int(2 * PT))))
    ours.smooth = False
    _end_label(ours, n - 1, ACCENT_TEXT, bold=True)

    bench.graphicalProperties = GraphicalProperties(
        ln=LineProperties(solidFill=CONTEXT, w=W_CONTEXT, prstDash="dash", cap="rnd"))
    bench.marker, bench.smooth = Marker(symbol="none"), False
    _end_label(bench, n - 1, CONTEXT, bold=True)

    y = _axis(chart.y_axis, GRID, TEXT_2)
    y.scaling.min, y.scaling.max, y.majorUnit = 0, 10, 2
    y.majorGridlines = _gridlines()
    y.title = _title("# OF INVENTORY TURNS", 8, TEXT_2, rot=-5400000)
    y.crossBetween = "midCat"

    x = _axis(chart.x_axis, GRID, TEXT_2)
    x.majorGridlines = None
    x.title = _title("FISCAL YEAR ENDING 12/31", 8, TEXT_2, left=0.175, top=0.875)
    x.lblOffset = 100

    year, caption = EVENT
    x_index = int(df.index[df["fiscal_year"] == year][0]) + 1      # derived, never typed
    chart += event_line(ws, ev_row, x_index, caption, n)           # combine LAST

    ws.add_chart(chart, f"B{head + 2}")

    src = ws.cell(last + 22, 2, SOURCE)
    src.font = Font(name=FONT_BODY, size=9, color=TEXT_2)


def _gridlines():
    from openpyxl.chart.axis import ChartLines
    return ChartLines(spPr=GraphicalProperties(
        ln=LineProperties(solidFill=GRID, w=int(0.9 * PT))))


# ---------------------------------------------------------------------------
df = pd.read_csv(CSV)

wb = Workbook()
# Rename BEFORE building: add_data() resolves a Reference to a string there and
# then, so a later rename leaves the chart pointing at "Sheet" -> #REF!.
wb.active.title = "replica"
build_replica(wb.active, df)
build_standup(wb.create_sheet("standup"), df)

wb.save(XLSX)
print(f"wrote {XLSX}")
