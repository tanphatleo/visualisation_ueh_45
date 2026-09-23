"""
The style guide, expressed for openpyxl — the Excel counterpart of style/finance_style.py.

Implements style/style_chart.md v1.0 for native Excel charts:
  * role colours, never a hex at the call site
  * the four-band layout (title 12% / subtitle 8% / visual 75% / source 5%) built
    from worksheet cells, because Excel charts have no native subtitle or source line
  * horizontal gridlines only, no chart border, no y-axis line, no tick marks
  * one accent per chart; everything else grey
"""

from openpyxl.chart.axis import ChartLines, DisplayUnitsLabelList
from openpyxl.chart.label import DataLabelList
from openpyxl.chart.marker import Marker
from openpyxl.chart.series import DataPoint
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.chart.text import RichText, Text
from openpyxl.chart.title import Title
from openpyxl.drawing.line import LineProperties
from openpyxl.drawing.text import (CharacterProperties, Font as DFont, Paragraph,
                                   ParagraphProperties, RegularTextRun, RichTextProperties)
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

# --------------------------------------------------------------------------
# Colour roles — style_chart.md §2.1. openpyxl wants RRGGBB with no '#'.
# --------------------------------------------------------------------------
SURFACE = "FFFFFF"
WASH = "F7F7F7"
GRID = "E8E8E8"
CONTEXT_FILL = "BDBDBD"
CONTEXT_LINE = "8C8C8C"
TEXT_2 = "595959"
INK = "262A33"
ACCENT = "0F5499"
ALERT = "990F3D"
GOOD = "0D7680"
BAD = "990F3D"
SERIES = ["262A33", "0F5499", "3D9199", "93B2D1"]   # §2.2, ordered by lightness
SERIES_4_FILL_ONLY = "93B2D1"

# Ordered scales — §2.6. Light = low, dark = high, always.
SEQUENTIAL = ["E6EEF6", "B8CFE4", "7FA6C9", "3F73A6", "0F5499", "0F3762"]
DIVERGING = ["990F3D", "C97C93", "F2F2F2", "7FB3B8", "0D7680"]

# Six branches, four series colours. §2.2 caps bar charts at 4 series and line
# charts at 3, and the palette is not to be extended — the guide's answer is
# small multiples, then group the tail, then highlight one against grey. Use
# branch_colours() for the third and ordered_colours() when the categories
# genuinely have an order (branches ranked by revenue, months, bins).
BRANCH_ORDER = ["HG-Q1", "HG-Q7", "HG-Q3", "HG-BT", "HG-TD", "HG-GV"]


def branch_colours(branches, highlight=None, bars=True):
    """Grey everything, accent the one the takeaway title is about."""
    base = CONTEXT_FILL if bars else CONTEXT_LINE
    return [ACCENT if b == highlight else base for b in branches]


def ordered_colours(n):
    """A sequential ramp for ordered categories. Never for an unordered set."""
    assert n <= len(SEQUENTIAL), (
        "%d steps asked of a 6-step ramp; group the tail instead" % n)
    step = (len(SEQUENTIAL) - 1) / max(1, n - 1) if n > 1 else 0
    return [SEQUENTIAL[min(len(SEQUENTIAL) - 1, int(round(i * step)))]
            for i in range(n)]

FONT = "Aptos"
# Chart canvas = the guide's 9.0in x 4.5in chart block, in centimetres.
CHART_W, CHART_H = 22.86, 11.43

# EMU line widths
PT = 12700
W_THIN, W_CONTEXT, W_HIGHLIGHT = int(0.75 * PT), int(1.5 * PT), int(2.5 * PT)

VND0 = '#,##0'
# Plain thousands separator, no currency symbol and no scaling suffix: the cell
# shows exactly the number it holds, in VND.
VND_M = '#,##0'
# Plain format scaled to millions, for DATA LABELS on charts that have no
# value axis to hang Excel's display units on. Still just digits.
MONEY_MILLIONS = '#,##0,,'
# One decimal, for label sets where whole millions would round the story away.
MONEY_MILLIONS_1 = '#,##0.0,,'
PCT0, PCT1 = '0%', '0.0%'


# --------------------------------------------------------------------------
# text helpers
# --------------------------------------------------------------------------
def _cp(size_pt, color, bold=False):
    return CharacterProperties(sz=int(size_pt * 100), b=bold, solidFill=color,
                               latin=DFont(typeface=FONT))


def _rich(size_pt, color, bold=False):
    """Text properties for axes and data labels."""
    cp = _cp(size_pt, color, bold)
    return RichText(bodyPr=RichTextProperties(),
                    p=[Paragraph(pPr=ParagraphProperties(defRPr=cp), endParaRPr=cp)])


def axis_title(axis, text, size_pt=10, color=TEXT_2):
    """
    An axis title that reads horizontally.

    Excel rotates a value-axis title 90° by default, which the guide forbids
    (§3: axis titles sit horizontally and left-aligned, never rotated). The
    rotation has to be forced off in the body properties — setting the title
    as a plain string gets you Excel's default every time.
    """
    cp = _cp(size_pt, color)
    para = Paragraph(pPr=ParagraphProperties(algn="l", defRPr=cp),
                     r=[RegularTextRun(rPr=cp, t=text)])
    axis.title = Title(
        tx=Text(rich=RichText(bodyPr=RichTextProperties(rot=0, vert="horz"),
                              p=[para])),
        overlay=False)
    return axis


def chart_message(message, descriptive=None, size_pt=12, sub_pt=10):
    """
    A two-line chart title: the takeaway, then the descriptive label.

    Rule 2 of the guide asks for a sentence stating the point; a reader still
    needs to know what is plotted and in what units. Both belong on the chart,
    because the moment it is copied into a deck the worksheet header does not
    go with it. The message is bold ink, the descriptive line is smaller and
    grey, so the eye takes them in that order.

    The source line is NOT part of this. It stays on the worksheet — a chart
    that carries its own provenance in the title has three titles.
    """
    cp = _cp(size_pt, INK, bold=True)
    paras = [Paragraph(pPr=ParagraphProperties(algn="l", defRPr=cp),
                       r=[RegularTextRun(rPr=cp, t=message)])]
    if descriptive:
        cs = _cp(sub_pt, TEXT_2)
        paras.append(Paragraph(pPr=ParagraphProperties(algn="l", defRPr=cs),
                               r=[RegularTextRun(rPr=cs, t=descriptive)]))
    return Title(tx=Text(rich=RichText(p=paras)), overlay=False)


def chart_title(text, size_pt=14, color=INK):
    cp = _cp(size_pt, color, bold=True)
    para = Paragraph(pPr=ParagraphProperties(defRPr=cp),
                     r=[RegularTextRun(rPr=cp, t=text)])
    return Title(tx=Text(rich=RichText(p=[para])), overlay=False)


# --------------------------------------------------------------------------
# worksheet furniture — the guide's four bands
# --------------------------------------------------------------------------
def as_text(value):
    """
    Force a string to be stored as TEXT.

    openpyxl treats any string starting with '=' as a formula. Excel cannot parse
    prose, strips the cell, and reports "Removed Records: Formula". Prefixing a
    zero-width space would change the text, so instead we let the caller write the
    value and then pin the cell's data type — see write_text_cell below.
    """
    return value


def write_text_cell(ws, row, col, value, font=None, alignment=None):
    """Write a value that must never be interpreted as a formula."""
    c = ws.cell(row, col, value)
    if isinstance(value, str) and value.startswith("="):
        c.data_type = "s"          # pin as string, whatever it looks like
    if font is not None:
        c.font = font
    if alignment is not None:
        c.alignment = alignment
    return c


def write_header(ws, row, title, subtitle, col=2):
    """Takeaway title + descriptive subtitle in cells (Excel has no chart subtitle)."""
    c = ws.cell(row, col, title)
    c.font = Font(name=FONT, size=16, bold=True, color=INK)
    c.alignment = Alignment(vertical="center")
    ws.row_dimensions[row].height = 24
    c = ws.cell(row + 1, col, subtitle)
    c.font = Font(name=FONT, size=11, color=TEXT_2)
    return row + 2


def write_source(ws, row, text, col=2):
    c = ws.cell(row, col, text)
    c.font = Font(name=FONT, size=9, color=CONTEXT_LINE)
    return row + 1


def write_table_header(ws, row, headers, col=2):
    for j, h in enumerate(headers):
        c = ws.cell(row, col + j, h)
        c.font = Font(name=FONT, size=11, bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=ACCENT)
        c.alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[row].height = 20
    return row + 1


def note(ws, row, text, col=2, color=TEXT_2, italic=True):
    c = ws.cell(row, col, text)
    if isinstance(text, str) and text.startswith("="):
        c.data_type = "s"          # prose, not a formula
    c.font = Font(name=FONT, size=10, color=color, italic=italic)
    return row + 1


def widths(ws, spec, start=2):
    for j, w in enumerate(spec):
        ws.column_dimensions[get_column_letter(start + j)].width = w


# --------------------------------------------------------------------------
# chart decluttering — §1 rule 6, §4
# --------------------------------------------------------------------------
def _axis_off(axis):
    axis.delete = False
    axis.majorTickMark = "none"
    axis.minorTickMark = "none"
    axis.txPr = _rich(11, TEXT_2)
    axis.spPr = GraphicalProperties(ln=LineProperties(noFill=True))


def declutter(chart, gridlines=True, x_line=True, y_labels=True, num_fmt=None,
              gap=40):
    """Strip Excel's default chart junk and apply the guide's frame."""
    chart.width, chart.height = CHART_W, CHART_H
    chart.roundedCorners = False
    chart.style = None

    # Excel's default gap between bars is 150% of the bar width, which spends
    # more of the plot area on background than on data. The guide calls for a
    # 40% gap (§4, BAR_WIDTH 0.72) — the bars get wider, the chart gets easier
    # to compare across, and nothing is lost.
    if gap is not None and hasattr(chart, "gapWidth"):
        chart.gapWidth = gap
    # no chart-area border, white plot area
    chart.graphical_properties = GraphicalProperties(
        solidFill=SURFACE, ln=LineProperties(noFill=True))

    # Pie and doughnut charts have no axes at all, so everything below is
    # skipped for them rather than guarded at every call site.
    if not hasattr(chart, "x_axis"):
        return chart

    for ax in (chart.x_axis, chart.y_axis):
        _axis_off(ax)
    # the x (category) baseline stays on column charts, per §4
    if x_line:
        chart.x_axis.spPr = GraphicalProperties(
            ln=LineProperties(solidFill=TEXT_2, w=W_THIN))

    chart.y_axis.majorGridlines = (
        ChartLines(spPr=GraphicalProperties(
            ln=LineProperties(solidFill=GRID, w=W_THIN))) if gridlines else None)
    chart.x_axis.majorGridlines = None

    if not y_labels:
        chart.y_axis.delete = True
    if num_fmt:
        chart.y_axis.numFmt = num_fmt
        chart.y_axis.number_format = num_fmt
    return chart


def display_units(chart, unit="millions"):
    """
    Excel's native axis Display Units. The axis shows 1.048 instead of
    1.048.000.000 while the cell keeps the full number. Name the unit in the
    axis title — we suppress Excel's own unit caption.
    unit: "thousands" | "millions" | "billions"
    """
    chart.y_axis.dispUnits = DisplayUnitsLabelList(builtInUnit=unit)
    return chart


def no_legend(chart):
    chart.legend = None
    return chart


def legend_bottom(chart):
    if chart.legend is not None:
        chart.legend.position = "b"
        chart.legend.overlay = False
        chart.legend.txPr = _rich(11, TEXT_2)
    return chart


def fill(series, color, line=False):
    """Solid fill for bars/areas, or a coloured line for line charts."""
    if line:
        series.graphicalProperties = GraphicalProperties(
            ln=LineProperties(solidFill=color, w=W_HIGHLIGHT))
        series.marker = Marker(symbol="none")
        series.smooth = False
    else:
        series.graphicalProperties = GraphicalProperties(
            solidFill=color, ln=LineProperties(noFill=True))
    return series


def context_line(series):
    series.graphicalProperties = GraphicalProperties(
        ln=LineProperties(solidFill=CONTEXT_LINE, w=W_CONTEXT))
    series.marker = Marker(symbol="none")
    series.smooth = False
    return series


def reference_line(series, colour=GRID):
    """A target or benchmark drawn as a series: thin, dashed, behind the data.

    Not a line a reader compares two values on — a line that says where the
    data ought to be. It is dashed because §2.2 reserves solid strokes for
    measured things, and grey because it must not compete with the series the
    title is about.
    """
    series.graphicalProperties = GraphicalProperties(
        ln=LineProperties(solidFill=colour, w=W_THIN, prstDash="dash"))
    series.marker = Marker(symbol="none")
    series.smooth = False
    return series


def invisible(series):
    """The hidden base of a waterfall, or the hidden half of a gauge."""
    series.graphicalProperties = GraphicalProperties(
        solidFill=None, noFill=True, ln=LineProperties(noFill=True))
    return series


def highlight_points(series, n, highlight_idx, base=CONTEXT_FILL, accent=ACCENT):
    """Grey every category except the one the takeaway title is about (§2.3)."""
    series.graphicalProperties = GraphicalProperties(
        solidFill=base, ln=LineProperties(noFill=True))
    series.data_points = [
        DataPoint(idx=i, spPr=GraphicalProperties(
            solidFill=(accent if i in highlight_idx else base),
            ln=LineProperties(noFill=True)))
        for i in range(n)]
    return series


def data_labels(chart, num_fmt=VND_M, size=11, color=TEXT_2, position=None):
    dl = DataLabelList(showVal=True, showSerName=False, showCatName=False,
                       showLegendKey=False, showPercent=False, showBubbleSize=False)
    # openpyxl writes <numFmt formatCode="..."/> with no sourceLinked attribute,
    # and Excel then defaults sourceLinked to TRUE and uses the cell's format
    # instead — silently ignoring this one. Passing a NumFmt object does not help
    # (the descriptor str()s it), so build_workbook.py patches the attribute into
    # the saved package. See fix_label_number_formats there.
    dl.numFmt = num_fmt
    dl.txPr = _rich(size, color)
    if position:
        dl.dLblPos = position
    chart.dataLabels = dl
    return chart


def event_lines(ws, row, events, n_categories, col=2, colour=ALERT):
    """
    Vertical event markers, built INTO the chart rather than drawn on top of it.

    A drawn shape does not move when the data updates, drifts when the chart is
    resized or pasted into a deck, and carries no text. So the line is a
    two-point scatter series on a hidden secondary axis instead — the technique
    from Storytelling with Data, "Tactical tip: embedding a vertical reference
    line in Excel", with two additions:

      * the secondary x axis is pinned to 0,5 .. n+0,5 so the line lands exactly
        on its period. The original leaves it to autoscale and loses precision.
      * the caption rides on the line: the series is NAMED after the event, and
        the top point alone shows a data label displaying the series name. Move
        the line and the words move with it.

    §1 rule 3 allows a secondary axis only where nothing else will do. This is
    that case, and it is the weakest possible form of it: the axis is hidden, it
    carries no series a reader compares against another, and it encodes nothing
    but the position of a line.

    events        [(x_index, caption, top)] — x_index is 1-based over the
                  categories; top is where the line stops, 1,0 = full height.
                  Drop the second of a close pair to about 0,88 so the captions
                  do not collide.
    returns       (scatter_chart, next_free_row)
    """
    from openpyxl.chart import Reference, ScatterChart, Series
    from openpyxl.chart.data_source import StrRef
    from openpyxl.chart.label import DataLabel
    from openpyxl.chart.series import SeriesLabel

    letter = get_column_letter(col)
    sc = ScatterChart()
    r = row
    for x_index, caption, top in events:
        c = ws.cell(r, col, caption)
        c.font = Font(name=FONT, size=9, color=TEXT_2, italic=True)
        for k, y in enumerate((0, top)):
            ws.cell(r + k, col + 1, x_index).font = Font(name=FONT, size=9, color=TEXT_2)
            ws.cell(r + k, col + 2, y).font = Font(name=FONT, size=9, color=TEXT_2)

        s = Series(Reference(ws, min_col=col + 2, min_row=r, max_row=r + 1),
                   Reference(ws, min_col=col + 1, min_row=r, max_row=r + 1))
        # Series(title=...) calls str() on whatever it is given, so the reference
        # has to be attached afterwards or the caption becomes a repr.
        s.tx = SeriesLabel(strRef=StrRef(f"'{ws.title}'!${letter}${r}"))
        s.graphicalProperties = GraphicalProperties(
            ln=LineProperties(solidFill=colour, w=W_THIN, prstDash="dash"))
        s.marker = Marker(symbol="none")

        off = dict(showVal=False, showCatName=False, showLegendKey=False,
                   showPercent=False, showBubbleSize=False)
        top_label = DataLabel(idx=1, showSerName=True, dLblPos="t", **off)
        top_label.txPr = _rich(9, colour)
        s.dLbls = DataLabelList(dLbl=[DataLabel(idx=0, showSerName=False, **off), top_label],
                                showSerName=False, **off)
        sc.series.append(s)
        r += 2

    # The secondary pair must sit OPPOSITE the primary pair and cross at max —
    # copied from what Excel writes for the same combo. Left at openpyxl's
    # defaults they land on top of the primary axes and Excel quietly rescales
    # the primary axis to fit the 0–1 helper values, flattening the real data.
    sc.x_axis.axId, sc.y_axis.axId = 300, 200
    sc.x_axis.crossAx, sc.y_axis.crossAx = 200, 300     # each must point at the other
    sc.x_axis.axPos, sc.y_axis.axPos = "t", "r"
    sc.x_axis.crosses = sc.y_axis.crosses = "max"
    sc.x_axis.scaling.min, sc.x_axis.scaling.max = 0.5, n_categories + 0.5
    sc.y_axis.scaling.min, sc.y_axis.scaling.max = 0, 1
    sc.x_axis.majorGridlines = sc.y_axis.majorGridlines = None
    # Hide the pair, but do NOT delete it: a deleted axis makes Excel reassign the
    # series to the primary axis, which then rescales to 0-1 and flattens the data.
    for ax in (sc.x_axis, sc.y_axis):
        ax.delete = False
        ax.majorTickMark = ax.minorTickMark = "none"
        # tickLblPos = "none" cannot be used: openpyxl reads the string "none" as
        # "unset" and drops the element, leaving the labels on. ";;;" is Excel's
        # own show-nothing number format and does the job.
        ax.numFmt = ";;;"
        ax.spPr = GraphicalProperties(ln=LineProperties(noFill=True))
    return sc, r


def label_last_point(series, n_points, colour, size=11, position="r"):
    """
    Name the series at the end of its own line, and drop the legend.

    Rule 6 of the guide — direct labelling beats legends — and on a combined
    chart it is the only thing that works reliably. openpyxl builds a combo by
    merging two chart objects, and the legend that survives is not always the
    one that was configured; event-line series then appear in it as data, which
    is exactly what they are not. Labelling the series itself sidesteps the
    whole question.
    """
    from openpyxl.chart.label import DataLabel

    off = dict(showVal=False, showCatName=False, showLegendKey=False,
               showPercent=False, showBubbleSize=False)
    lbl = DataLabel(idx=n_points - 1, showSerName=True, dLblPos=position, **off)
    lbl.txPr = _rich(size, colour, bold=True)
    series.dLbls = DataLabelList(dLbl=[lbl], showSerName=False, **off)
    return series


def label_point(series, idx, colour, size=11, position="t"):
    """
    Name the series at a chosen point rather than at its last one.

    `label_last_point` is the default and it is right almost always — but two
    series that converge share their last point, and two labels at the same
    coordinates is a collision the reader has to untangle. Where the series
    separate is both legible and more informative: it is the place the chart
    is making its argument.
    """
    from openpyxl.chart.label import DataLabel

    off = dict(showVal=False, showCatName=False, showLegendKey=False,
               showPercent=False, showBubbleSize=False)
    lbl = DataLabel(idx=idx, showSerName=True, dLblPos=position, **off)
    lbl.txPr = _rich(size, colour, bold=True)
    series.dLbls = DataLabelList(dLbl=[lbl], showSerName=False, **off)
    return series


def hide_legend_entries(chart, indices):
    """
    Drop specific series from the legend without removing them from the chart.

    Event lines carry their caption on the line itself, so a legend entry for
    them is pure repetition — and it reads as data, which they are not.
    """
    from openpyxl.chart.legend import LegendEntry
    if chart.legend is not None:
        chart.legend.legendEntry = [LegendEntry(idx=i, delete=True) for i in indices]
    return chart
